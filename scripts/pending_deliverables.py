#!/usr/bin/env python3
"""List deliverables the dedup log promises that do not yet exist on disk.

The scan pass assigns IDs (``bulletin:CTI26-07``, ``hunt:TH26-03``, an AWR id for
today) in ``state/dedup-log.md``. This driver turns those promises into a work
batch for the builder pass: for each pending ID it emits the log entry that
carries the evidence, so the builder never has to research from scratch.

Exit 0 and write the batch when something is pending; exit 3 when nothing is.
"""
from __future__ import annotations

import argparse
import datetime as dt
import json
import os
import re
import sys
from pathlib import Path

OUT = Path(os.environ.get("OUTPUT_DIR", "out")).resolve()
DEDUP = OUT / "state" / "dedup-log.md"
BULLETIN = re.compile(r"bulletin:\s*(CTI\d{2}-\d{2})")
HUNT = re.compile(r"hunt:\s*(TH\d{2}-\d{2})")
AWR = re.compile(r"\b(AWR-\d{4}-\d{2}-\d{2})\b")


# A file whose name carries the ID is not enough: a builder fed a spec with the wrong
# keys emits a hollow document with placeholder text, and that must be rebuilt.
MIN_PARAS = {"bulletin": 15, "hunt": 30, "awareness": 8}
PLACEHOLDERS = ("TH-26-XX", "TH-26-NN", "CTIYY-NN")
# Sections April removed from the package format on 2026-08-24. Their presence means a
# reverted builder or a spec still carrying the old keys, so the document is rebuilt.
REMOVED_SECTIONS = {"hunt": ("Findings Classification Key", "Coverage & Gaps", "Execution Log",
                             "Outcome & Classification", "Overall Finding", "Pending Review")}
# Where task-build.md tells the builder to save each kind, from config/paths.json. Used only
# to report a document saved somewhere else; it is still counted as built.
CANONICAL = {"bulletin": "bulletins", "hunt": "hunts", "awareness": "employee-posts"}


def shell_reason(path: Path, ident: str, kind: str) -> str | None:
    """Return why the document is a shell, or None when it looks complete."""
    try:
        import docx  # python-docx, present in the image
        d = docx.Document(str(path))
        paras = [p.text for p in d.paragraphs if p.text.strip()]
        text = "\n".join(paras) + "\n".join(c.text for t in d.tables for r in t.rows for c in r.cells)
        # the awareness post carries its ID in the footer only
        text += "\n".join(p.text for sec in d.sections for part in (sec.header, sec.footer) for p in part.paragraphs)
    except Exception as exc:  # unreadable file counts as missing
        return f"unreadable ({exc.__class__.__name__})"
    if ident not in text:
        return f"ID {ident} not in document text"
    if len(paras) < MIN_PARAS.get(kind, 8):
        return f"only {len(paras)} paragraphs"
    hit = next((ph for ph in PLACEHOLDERS if ph in text), None)
    if hit:
        return f"placeholder text {hit!r}"
    gone = next((sec for sec in REMOVED_SECTIONS.get(kind, ()) if sec in text), None)
    return f"carries removed section {gone!r} (reverted builder or stale spec)" if gone else None


def on_disk(pattern: str, kind: str) -> tuple[set[str], dict[str, str], dict[str, str]]:
    """IDs with a complete document, IDs whose document is a shell, and IDs saved off path.

    The whole output tree is searched apart from ``state`` (scratch and ledgers). Searching
    only the canonical folder made a builder that saved elsewhere look like it had built
    nothing, so the ID was queued again in every pass and burned a batch slot each time.
    A document found off path still counts; the stray path is reported so the drift is
    visible instead of silent.
    """
    ids: set[str] = set(); shells: dict[str, str] = {}; strays: dict[str, str] = {}
    rx = re.compile(pattern)
    for f in OUT.rglob("*.docx"):
        rel = f.relative_to(OUT)
        if rel.parts[0] == "state":
            continue
        m = rx.search(f.name)
        if not m:
            continue
        why = shell_reason(f, m.group(0), kind)
        if why:
            shells.setdefault(m.group(0), f"{rel}: {why}")
            continue
        ids.add(m.group(0))
        if not str(rel).startswith(CANONICAL[kind]):
            strays[m.group(0)] = str(rel)
    # a complete document anywhere beats a shell somewhere else
    return ids, {k: v for k, v in shells.items() if k not in ids}, strays


def blocked_reason(ident: str) -> str | None:
    """The builder's own 'cannot build' marker for this run, or None.

    task-build.md tells the builder to write ``state/_work/<ID>.failed.md`` when an item
    cannot be built at all. Honouring it stops that item consuming a slot in every
    remaining pass. The entrypoint clears the markers before the loop, so the block lasts
    one run and the next run retries from scratch.
    """
    marker = OUT / "state" / "_work" / f"{ident}.failed.md"
    if not marker.exists():
        return None
    first = next((ln.strip() for ln in marker.read_text().splitlines() if ln.strip()), "")
    return first[:200] or "no reason given"


def entries(text: str) -> list[str]:
    """Split the log into entry blocks: a dated header line plus its indented lines."""
    blocks: list[str] = []
    current: list[str] = []
    for line in text.splitlines():
        if re.match(r"^\d{4}-\d{2}-\d{2} \| ", line):
            if current:
                blocks.append("\n".join(current))
            current = [line]
        elif current and (line.startswith("  ") or not line.strip()):
            current.append(line)
        elif current:
            blocks.append("\n".join(current)); current = []
    if current:
        blocks.append("\n".join(current))
    return blocks


def main() -> int:
    ap = argparse.ArgumentParser(description=__doc__)
    ap.add_argument("--batch", type=int, default=4, help="max items per builder pass")
    ap.add_argument("--out", type=Path, default=OUT / "state" / "_work" / "pending.json")
    ap.add_argument("--today", default=dt.date.today().isoformat())
    args = ap.parse_args()
    if not DEDUP.exists():
        print("[pending] no dedup log; nothing to build"); return 3

    text = DEDUP.read_text()
    have_b, shell_b, stray_b = on_disk(r"CTI\d{2}-\d{2}", "bulletin")
    have_h, shell_h, stray_h = on_disk(r"TH\d{2}-\d{2}", "hunt")
    have_a, shell_a, stray_a = on_disk(r"AWR-\d{4}-\d{2}-\d{2}", "awareness")
    shells = {**shell_b, **shell_h, **shell_a}
    for ident, where in {**stray_b, **stray_h, **stray_a}.items():
        print(f"[pending] {ident} built outside its folder: {where} (counted, but fix the save path)")

    items: list[dict] = []
    seen: set[str] = set()
    for block in entries(text):
        for rx, kind, have in ((BULLETIN, "bulletin", have_b), (HUNT, "hunt", have_h)):
            for ident in rx.findall(block):
                if ident not in have and ident not in seen:
                    seen.add(ident)
                    item = {"id": ident, "kind": kind, "evidence": block.strip()}
                    if ident in shells:
                        item["rebuild"] = f"existing file is a shell, overwrite it: {shells[ident]}"
                    items.append(item)
    awr = f"AWR-{args.today}"
    if awr in text and awr not in have_a:
        item = {"id": awr, "kind": "awareness", "evidence": "topic chosen in today's scan report and dedup log"}
        if awr in shells:
            item["rebuild"] = f"existing file is a shell, overwrite it: {shells[awr]}"
        items.append(item)
    brief = f"CVE-VM-{args.today}"
    if not (OUT / "registers" / f"{brief}.docx").exists() and (OUT / "registers" / "CVE_Register_2026.xlsx").exists():
        items.append({"id": brief, "kind": "cve_brief", "evidence": "build from registers/CVE_Register_2026.xlsx, Latest Run tab"})

    blocked = {i["id"]: r for i in items if (r := blocked_reason(i["id"]))}
    items = [i for i in items if i["id"] not in blocked]
    remaining = len(items)
    for ident, why in blocked.items():
        print(f"[pending] {ident} skipped this run, builder could not build it: {why}")
    # bulletins and hunts carry the most work; brief and awareness are cheap, keep them in the first batch
    order = {"cve_brief": 0, "awareness": 1, "bulletin": 2, "hunt": 3}
    items.sort(key=lambda i: order[i["kind"]])
    # a hunt package is roughly two bulletins' worth of work: weight it so a batch
    # never asks one context for more than two of them
    batch: list[dict] = []; weight = 0
    for item in items:
        w = 2 if item["kind"] == "hunt" else 1
        if batch and weight + w > args.batch:
            break
        batch.append(item); weight += w
    summary = ", ".join(f"{i['kind']}={sum(1 for x in items if x['kind']==i['kind'])}" for i in {x["kind"]: x for x in items}.values())
    if not batch:
        tail = f"; {len(blocked)} blocked this run: {sorted(blocked)}" if blocked else ""
        print(f"[pending] nothing left to build: every ID in the dedup log has a file on disk{tail}")
        return 3
    args.out.parent.mkdir(parents=True, exist_ok=True)
    args.out.write_text(json.dumps({"today": args.today, "remaining": remaining, "batch": batch}, indent=2))
    rebuilds = [i["id"] for i in items if i.get("rebuild")]
    print(f"[pending] {remaining} pending ({summary}); shells to rebuild: {rebuilds or 'none'}; "
          f"this batch: {[i['id'] for i in batch]} -> {args.out}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
