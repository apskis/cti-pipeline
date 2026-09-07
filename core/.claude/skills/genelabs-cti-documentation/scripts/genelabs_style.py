"""
Shared GeneLabs CTI house style for official program documentation.

Palette, typography and page furniture are identical to genelabs-cti-bulletin
and genelabs-threat-hunt-report so that every artifact the CTI program ships
reads as one system. Import from here; never hand pick colours.
"""
import os
import re
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Emu, Inches, Pt, RGBColor

# ---- corporate palette (shared across all CTI artifacts) -------------------
ORANGE       = "FF7109"   # GeneLabs orange: rules, accents
ORANGE_DEEP  = "E05F00"   # deepened orange: small heading text
ABBEY        = "444648"   # GeneLabs Abbey: body and headline text
GREY         = "676765"   # corporate grey: meta text, wordmark, footer
RULE         = "D9D9D9"   # light grey rules; stays visible on Word dark canvas
LABEL_FILL   = "F4F4F4"   # table label column
CALLOUT_FILL = "FFF4E8"   # callout and table header fill
LINK_BLUE    = "1565C0"
FONT         = "Arial"

TLP_COLORS = {
    "RED": "B71C1C", "AMBER+STRICT": "EF6C00", "AMBER": "EF6C00",
    "GREEN": "2E7D32", "CLEAR": "1565C0", "WHITE": "1565C0",
}
DEFAULT_TLP = "AMBER+STRICT"

HERE = os.path.dirname(os.path.abspath(__file__))
ASSETS = os.path.join(HERE, "..", "assets")
DEFAULT_WORDMARK = os.path.join(ASSETS, "genelabs_wordmark_grey.png")
DEFAULT_BANNER = os.path.join(ASSETS, "cti_banner.jpg")
DEFAULT_TEMPLATE = os.path.join(ASSETS, "cti_doc_template.docx")

BANNER_ASPECT = 274.0 / 2550.0   # height / width of the house banner


# ---- run and paragraph primitives -----------------------------------------
def run(p, text, size=10, bold=False, italic=False, color=ABBEY, caps=False,
        underline=False, spacing=None):
    r = p.add_run(text)
    r.font.name = FONT
    r.font.size = Pt(size)
    r.font.bold = bold
    r.font.italic = italic
    r.font.underline = underline
    r.font.color.rgb = RGBColor.from_string(color)
    rpr = r._element.get_or_add_rPr()
    rf = rpr.find(qn('w:rFonts'))
    if rf is None:
        rf = OxmlElement('w:rFonts')
        rpr.insert(0, rf)
    for a in ('w:ascii', 'w:hAnsi', 'w:cs', 'w:eastAsia'):
        rf.set(qn(a), FONT)
    if caps:
        el = OxmlElement('w:caps')
        el.set(qn('w:val'), '1')
        rpr.append(el)
    if spacing:
        el = OxmlElement('w:spacing')
        el.set(qn('w:val'), str(spacing))
        rpr.append(el)
    return r


def para(doc, text="", size=10, color=ABBEY, bold=False, italic=False,
         space_before=0, space_after=6, align=None):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(space_before)
    p.paragraph_format.space_after = Pt(space_after)
    if align is not None:
        p.alignment = align
    if text:
        run(p, text, size=size, color=color, bold=bold, italic=italic)
    return p


def border(p, edge, color, sz=6, space=1):
    pPr = p._p.get_or_add_pPr()
    bdr = pPr.find(qn('w:pBdr'))
    if bdr is None:
        bdr = OxmlElement('w:pBdr')
        pPr.append(bdr)
    el = bdr.find(qn('w:' + edge))
    if el is None:
        el = OxmlElement('w:' + edge)
        bdr.append(el)
    el.set(qn('w:val'), 'single')
    el.set(qn('w:sz'), str(sz))
    el.set(qn('w:space'), str(space))
    el.set(qn('w:color'), color)


def shade_paragraph(p, fill):
    pPr = p._p.get_or_add_pPr()
    sh = OxmlElement('w:shd')
    sh.set(qn('w:val'), 'clear')
    sh.set(qn('w:fill'), fill)
    pPr.append(sh)


def shade_cell(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr()
    sh = OxmlElement('w:shd')
    sh.set(qn('w:val'), 'clear')
    sh.set(qn('w:fill'), fill)
    tcPr.append(sh)


def table_borders(table, color=RULE, sz=6):
    tbl = table._tbl
    tblPr = tbl.tblPr
    old = tblPr.find(qn('w:tblBorders'))
    if old is not None:
        tblPr.remove(old)
    borders = OxmlElement('w:tblBorders')
    for edge in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
        el = OxmlElement('w:' + edge)
        el.set(qn('w:val'), 'single')
        el.set(qn('w:sz'), str(sz))
        el.set(qn('w:space'), '0')
        el.set(qn('w:color'), color)
        borders.append(el)
    tblPr.append(borders)


def keep_with_next(p):
    pPr = p._p.get_or_add_pPr()
    el = OxmlElement('w:keepNext')
    el.set(qn('w:val'), '1')
    pPr.append(el)


def field(p, instr, size=8, color=GREY):
    r = p.add_run()
    r.font.name = FONT
    r.font.size = Pt(size)
    r.font.color.rgb = RGBColor.from_string(color)
    fld = OxmlElement('w:fldChar')
    fld.set(qn('w:fldCharType'), 'begin')
    r._r.append(fld)
    it = OxmlElement('w:instrText')
    it.set(qn('xml:space'), 'preserve')
    it.text = instr
    r._r.append(it)
    end = OxmlElement('w:fldChar')
    end.set(qn('w:fldCharType'), 'end')
    r._r.append(end)


def bullet(doc, text, size=10, color=ABBEY, indent=0.28):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.left_indent = Inches(indent)
    p.paragraph_format.first_line_indent = Inches(-indent)
    run(p, "▪  ", size=9, bold=True, color=ORANGE)
    run(p, text, size=size, color=color)
    return p


# ---- house furniture -------------------------------------------------------
def set_base_style(doc):
    st = doc.styles['Normal']
    st.font.name = FONT
    st.font.size = Pt(10)
    st.font.color.rgb = RGBColor.from_string(ABBEY)
    rpr = st.element.get_or_add_rPr()
    rf = rpr.find(qn('w:rFonts'))
    if rf is None:
        rf = OxmlElement('w:rFonts')
        rpr.insert(0, rf)
    for a in ('w:ascii', 'w:hAnsi', 'w:cs', 'w:eastAsia'):
        rf.set(qn(a), FONT)


def style_headings(doc):
    """Give Word's built in heading styles the CTI house look."""
    specs = {
        'Title':     dict(size=20, color=ABBEY,       bold=True,  caps=False, rule=True),
        'Heading 1': dict(size=11, color=ORANGE_DEEP, bold=True,  caps=True,  rule=True),
        'Heading 2': dict(size=10.5, color=ORANGE_DEEP, bold=True, caps=False, rule=False),
        'Heading 3': dict(size=10, color=ABBEY,       bold=True,  caps=False, rule=False),
        'Heading 4': dict(size=10, color=GREY,        bold=True,  caps=False, rule=False),
    }
    for name, s in specs.items():
        try:
            st = doc.styles[name]
        except KeyError:
            continue
        st.font.name = FONT
        st.font.size = Pt(s['size'])
        st.font.bold = s['bold']
        st.font.color.rgb = RGBColor.from_string(s['color'])
        rpr = st.element.get_or_add_rPr()
        rf = rpr.find(qn('w:rFonts'))
        if rf is None:
            rf = OxmlElement('w:rFonts')
            rpr.insert(0, rf)
        for a in ('w:ascii', 'w:hAnsi', 'w:cs', 'w:eastAsia'):
            rf.set(qn(a), FONT)
        for tag in ('w:caps', 'w:spacing'):
            old = rpr.find(qn(tag))
            if old is not None:
                rpr.remove(old)
        if s['caps']:
            el = OxmlElement('w:caps')
            el.set(qn('w:val'), '1')
            rpr.append(el)
            sp = OxmlElement('w:spacing')
            sp.set(qn('w:val'), '20')
            rpr.append(sp)
        ppr = st.element.get_or_add_pPr()
        old = ppr.find(qn('w:pBdr'))
        if old is not None:
            ppr.remove(old)
        if s['rule']:
            bdr = OxmlElement('w:pBdr')
            bottom = OxmlElement('w:bottom')
            bottom.set(qn('w:val'), 'single')
            bottom.set(qn('w:sz'), '12')
            bottom.set(qn('w:space'), '3')
            bottom.set(qn('w:color'), ORANGE)
            bdr.append(bottom)
            ppr.append(bdr)
        st.paragraph_format.space_before = Pt(14 if name in ('Heading 1', 'Title') else 10)
        st.paragraph_format.space_after = Pt(4)
        st.paragraph_format.keep_with_next = True


def add_banner(doc, banner_path=None, min_top_margin_in=1.15):
    """Put the house banner in the page header of every section."""
    banner_path = banner_path or DEFAULT_BANNER
    if not os.path.exists(banner_path):
        return False
    for section in doc.sections:
        width = section.page_width - section.left_margin - section.right_margin
        header = section.header
        header.is_linked_to_previous = False
        # Clear whatever page furniture the document carried before (logo
        # blocks, title tables, page number lines) so the house banner is the
        # only thing at the top. Content controls are left alone: a watermark
        # lives in one, and removing it is a separate, explicit decision.
        for child in list(header._element):
            if child.tag != qn('w:sdt'):
                header._element.remove(child)
        p = header.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p.paragraph_format.space_after = Pt(0)
        p.add_run().add_picture(banner_path, width=Emu(int(width)))
        banner_h = Emu(int(width * BANNER_ASPECT))
        needed = banner_h + Emu(int(0.35 * 914400))
        if section.top_margin < needed:
            section.top_margin = needed
        if section.header_distance > Emu(int(0.4 * 914400)):
            section.header_distance = Emu(int(0.4 * 914400))
    return True


def build_footer(doc, doc_id, doc_marking, tlp=DEFAULT_TLP,
                 copyright_line="", wordmark_path=None):
    """The shared CTI footer: wordmark, TLP marking, DOC #, live page fields."""
    wordmark_path = wordmark_path or DEFAULT_WORDMARK
    for section in doc.sections:
        footer = section.footer
        footer.is_linked_to_previous = False
        for p in list(footer.paragraphs):
            p._element.getparent().remove(p._element)
        for t in list(footer.tables):
            t._tbl.getparent().remove(t._tbl)

        rule = footer.add_paragraph()
        rule.paragraph_format.space_before = Pt(0)
        rule.paragraph_format.space_after = Pt(4)
        border(rule, 'top', ORANGE, sz=12, space=1)

        width = section.page_width - section.left_margin - section.right_margin
        table = footer.add_table(rows=1, cols=3, width=Emu(int(width)))
        table.autofit = False
        table.alignment = WD_TABLE_ALIGNMENT.LEFT
        fracs = (0.15, 0.58, 0.27)
        tblPr = table._tbl.tblPr
        layout = OxmlElement('w:tblLayout')
        layout.set(qn('w:type'), 'fixed')
        tblPr.append(layout)
        for row in table.rows:
            for cell, f in zip(row.cells, fracs):
                cell.width = Emu(int(width * f))
        cells = table.rows[0].cells

        p0 = cells[0].paragraphs[0]
        p0.paragraph_format.space_after = Pt(0)
        if os.path.exists(wordmark_path):
            p0.add_run().add_picture(wordmark_path, width=Inches(0.95))

        p1 = cells[1].paragraphs[0]
        p1.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p1.paragraph_format.space_after = Pt(0)
        run(p1, "TLP: ", size=8, bold=True, italic=True, color=ABBEY)
        run(p1, tlp, size=8, bold=True, italic=True,
            color=TLP_COLORS.get(tlp.upper(), ORANGE_DEEP))
        run(p1, "  |  " + doc_marking, size=8, color=GREY)
        if copyright_line:
            p1b = cells[1].add_paragraph()
            p1b.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p1b.paragraph_format.space_after = Pt(0)
            run(p1b, copyright_line, size=7, color=GREY)

        p2 = cells[2].paragraphs[0]
        p2.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        p2.paragraph_format.space_after = Pt(0)
        run(p2, "DOC # ", size=8, bold=True, color=ABBEY)
        run(p2, doc_id, size=8, color=GREY)
        p2b = cells[2].add_paragraph()
        p2b.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        p2b.paragraph_format.space_after = Pt(0)
        run(p2b, "Page ", size=8, color=GREY)
        field(p2b, " PAGE ")
        run(p2b, " of ", size=8, color=GREY)
        field(p2b, " NUMPAGES ")


def strip_watermarks(doc):
    """Remove stale DRAFT style watermarks stored as header content controls."""
    removed = 0
    for section in doc.sections:
        for hdr in (section.header, section.first_page_header,
                    section.even_page_header):
            if hdr is None:
                continue
            el = hdr._element
            from lxml import etree
            for sdt in el.findall(qn('w:sdt')):
                if b'Watermarks' in etree.tostring(sdt):
                    el.remove(sdt)
                    removed += 1
            for pict in list(el.iter(qn('w:pict'))):
                parent = pict.getparent()
                if parent is not None and b'textpath' in etree.tostring(pict).lower():
                    parent.remove(pict)
                    removed += 1
    return removed


def promote_headings(doc):
    """Shift heading levels so the document's top level sections are Heading 1.

    Documents drafted with Heading 3 as their top level read as subordinate
    once the house heading treatment is applied, and their outline is wrong for
    navigation and tables of contents. This fixes the level, never the words.
    """
    used = set()
    for para_ in doc.paragraphs:
        m = re.match(r"Heading (\d)$", para_.style.name)
        if m:
            used.add(int(m.group(1)))
    if not used:
        return 0
    shift = min(used) - 1
    if shift <= 0:
        return 0
    for para_ in doc.paragraphs:
        m = re.match(r"Heading (\d)$", para_.style.name)
        if m:
            new = max(1, int(m.group(1)) - shift)
            para_.style = doc.styles["Heading %d" % new]
    return shift


def set_grid(table, fractions, total_width):
    """Fix a table's column widths to fractions of the content width."""
    tblPr = table._tbl.tblPr
    old = tblPr.find(qn('w:tblLayout'))
    if old is not None:
        tblPr.remove(old)
    layout = OxmlElement('w:tblLayout')
    layout.set(qn('w:type'), 'fixed')
    tblPr.append(layout)
    table.autofit = False

    widths = [int(total_width * f) for f in fractions]
    grid = table._tbl.find(qn('w:tblGrid'))
    if grid is not None:
        table._tbl.remove(grid)
    grid = OxmlElement('w:tblGrid')
    for w in widths:
        gc = OxmlElement('w:gridCol')
        gc.set(qn('w:w'), str(int(w / 635)))   # EMU to twips
        grid.append(gc)
    tblPr.addnext(grid)

    old_w = tblPr.find(qn('w:tblW'))
    if old_w is not None:
        tblPr.remove(old_w)
    tw = OxmlElement('w:tblW')
    tw.set(qn('w:w'), str(int(sum(widths) / 635)))
    tw.set(qn('w:type'), 'dxa')
    tblPr.append(tw)

    for row in table.rows:
        for cell, w in zip(row.cells, widths):
            cell.width = Emu(w)


CHANGE_HISTORY_HEADERS = {"version", "date", "originator", "description of change"}


def fit_change_history(doc, fractions=(0.12, 0.15, 0.18, 0.55)):
    """Give Release History and Revision History tables usable column widths.

    Left to autofit, the Description of Change column collapses and the entry
    becomes unreadable, which defeats the point of keeping a change history.
    """
    section = doc.sections[0]
    width = section.page_width - section.left_margin - section.right_margin
    fixed = 0
    for t in doc.tables:
        if not t.rows or len(t.columns) != len(fractions):
            continue
        head = {c.text.strip().lower() for c in t.rows[0].cells}
        if CHANGE_HISTORY_HEADERS.issubset(head):
            set_grid(t, fractions, width)
            fixed += 1
    return fixed
