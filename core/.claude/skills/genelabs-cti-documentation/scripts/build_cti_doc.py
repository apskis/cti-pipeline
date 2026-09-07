"""
Build a new GeneLabs CTI program document (.docx) from a JSON spec.

    python build_cti_doc.py spec.json out.docx [template.docx]

For plans, procedures, policies, playbooks, methodologies and standards — the
governance layer of the CTI program. For a bulletin use genelabs-cti-bulletin;
for a hunt use genelabs-threat-hunt-package or genelabs-threat-hunt-report.
"""
import json
import os
import sys

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from genelabs_style import (  # noqa: E402
    ABBEY, CALLOUT_FILL, DEFAULT_TEMPLATE, DEFAULT_TLP, GREY, LABEL_FILL,
    ORANGE, ORANGE_DEEP, RULE, add_banner, border, build_footer, bullet,
    keep_with_next, para, run, set_base_style, shade_cell, shade_paragraph,
    style_headings, table_borders,
)

HEADING_STYLES = {1: "Heading 1", 2: "Heading 2", 3: "Heading 3", 4: "Heading 4"}


def add_doc_control(doc, rows):
    """The two column control block that opens every governed CTI document."""
    t = doc.add_table(rows=len(rows), cols=2)
    t.autofit = True
    for i, pair in enumerate(rows):
        label, value = (list(pair) + [""])[:2]
        lc, vc = t.rows[i].cells
        lc.text = ""
        vc.text = ""
        run(lc.paragraphs[0], label, size=9, bold=True, color=ABBEY)
        run(vc.paragraphs[0], value, size=9, color=GREY)
        shade_cell(lc, LABEL_FILL)
    table_borders(t, color=RULE)
    return t


def add_title_block(doc, title, subtitle=None, meta=None):
    p = para(doc, space_before=10, space_after=2)
    run(p, title, size=18, bold=True, color=ABBEY)
    border(p, 'bottom', ORANGE, sz=18, space=4)
    if subtitle:
        para(doc, subtitle, size=10, color=GREY, italic=True, space_after=4)
    if meta:
        p = para(doc, space_after=10)
        run(p, meta, size=8, color=GREY, caps=True, spacing=20)


def add_callout(doc, label, text):
    p = para(doc, space_before=8, space_after=8)
    shade_paragraph(p, CALLOUT_FILL)
    border(p, 'left', ORANGE, sz=24, space=6)
    run(p, label + "  ", size=10, bold=True, color=ABBEY)
    run(p, text, size=10, color=ABBEY)
    return p


def add_table(doc, spec):
    header = spec.get("header") or []
    rows = spec.get("rows") or []
    if not rows and not header:
        return None
    ncols = len(header) if header else max(len(r) for r in rows)
    t = doc.add_table(rows=len(rows) + (1 if header else 0), cols=ncols)
    t.autofit = True
    if header:
        for ci, cell_text in enumerate(header):
            cell = t.rows[0].cells[ci]
            cell.text = ""
            run(cell.paragraphs[0], str(cell_text), size=9, bold=True, color=ABBEY)
            shade_cell(cell, CALLOUT_FILL)
    off = 1 if header else 0
    for ri, row in enumerate(rows):
        for ci in range(ncols):
            cell = t.rows[ri + off].cells[ci]
            cell.text = ""
            value = str(row[ci]) if ci < len(row) else ""
            run(cell.paragraphs[0], value, size=9, color=ABBEY)
    table_borders(t, color=RULE)
    if spec.get("caption"):
        para(doc, spec["caption"], size=8, color=GREY, italic=True,
             space_before=2, space_after=8)
    return t


def add_section(doc, sec, level=1):
    heading = sec.get("heading")
    if heading:
        p = doc.add_paragraph(heading, style=HEADING_STYLES.get(level, "Heading 4"))
        keep_with_next(p)
    if sec.get("intro"):
        para(doc, sec["intro"], size=10, color=ABBEY, space_after=6)
    for text in sec.get("paragraphs", []):
        para(doc, text, size=10, color=ABBEY, space_after=6)
    for item in sec.get("bullets", []):
        bullet(doc, item)
    if sec.get("callout"):
        c = sec["callout"]
        add_callout(doc, c.get("label", "Note:"), c.get("text", ""))
    if sec.get("table"):
        add_table(doc, sec["table"])
    for sub in sec.get("subsections", []):
        add_section(doc, sub, level + 1)


def add_release_history(doc, spec):
    heading = spec.get("heading", "Release History")
    p = doc.add_paragraph(heading, style="Heading 1")
    keep_with_next(p)
    if spec.get("intro"):
        para(doc, spec["intro"], size=9, color=GREY, italic=True, space_after=4)
    add_table(doc, {
        "header": spec.get("columns",
                           ["Version", "Date", "Originator", "Description of Change"]),
        "rows": spec.get("rows", []),
    })


def build(spec, out_path, template_path=None):
    template_path = template_path or DEFAULT_TEMPLATE
    doc = Document(template_path) if os.path.exists(template_path) else Document()
    set_base_style(doc)
    style_headings(doc)

    if spec.get("doc_control"):
        add_doc_control(doc, spec["doc_control"])
    add_title_block(doc, spec["title"], spec.get("subtitle"), spec.get("meta"))
    if spec.get("summary"):
        add_callout(doc, spec.get("summary_label", "In brief:"), spec["summary"])
    for sec in spec.get("sections", []):
        add_section(doc, sec, sec.get("level", 1))
    if spec.get("release_history"):
        add_release_history(doc, spec["release_history"])
    for sec in spec.get("appendix", []):
        add_section(doc, sec, sec.get("level", 1))

    add_banner(doc)
    build_footer(doc, spec.get("doc_id", ""),
                 spec.get("marking", "CTI Program Documentation"),
                 tlp=spec.get("tlp", DEFAULT_TLP),
                 copyright_line=spec.get("copyright_line", ""))
    cp = doc.core_properties
    cp.title = spec["title"]
    cp.comments = ("GeneLabs Enterprise Information Security | "
                   "Cyber Threat Intelligence")
    doc.save(out_path)
    return out_path


if __name__ == "__main__":
    if len(sys.argv) < 3:
        print(__doc__)
        sys.exit(1)
    spec = json.load(open(sys.argv[1], encoding="utf-8"))
    template = sys.argv[3] if len(sys.argv) > 3 else None
    print("wrote", build(spec, sys.argv[2], template))
