#!/usr/bin/env python3
"""
Build an GeneLabs-branded CTI bulletin .docx from a JSON content spec, using the
bundled template (assets/bulletin_template.docx) as the branding anchor.

The template carries the GeneLabs "Enterprise Information Security / Cyber
Threat Intelligence" banner in its document header, so cloning it preserves the
house look. This script writes the body in the GeneLabs corporate identity
(GeneLabs orange accents, Abbey body text, corporate grey meta), renders inline
citations as superscript, hyperlinks the Sources list, and rebuilds the page
footer with the genelabs wordmark, TLP marking, DOC # field, live page numbers
and the copyright line.

Section flow follows the published house order (see CTI-26-07):
    report ID + date -> category tag -> title -> audience -> relevance callout
    -> detail table (no heading) -> Why This Matters -> What Happened
    -> Red Flags to Watch -> Recommended Actions -> Abbreviations -> Sources
    -> contact line

Usage:
    python build_bulletin.py content.json out.docx [path/to/bulletin_template.docx]

The wordmark asset (assets/genelabs_wordmark_grey.png) is resolved relative to
this script; pass a fourth argument to override it.

content.json schema (all keys optional except report_id, date, category, title,
severity, relevance, what_happened, sources):
{
  "report_id": "CTI-26-11",
  "date": "August 17, 2026",
  "category": "AWARENESS BULLETIN",       # must be one of CATEGORY_COLORS below
  "severity": "MEDIUM",                    # must be one of SEVERITY_COLORS below
  "title": "Bulletin title",
  "audience": "All Employees, SOC, Leadership",
  "supersedes": "CTI-26-07",               # optional
  "relevance": "Prose for the Relevance to GeneLabs callout (no prefix needed).",
  "detail_table": [["Threat Actor", "..."], ["MITRE ATT&CK", "T1566"]],
  "why_this_matters": ["bullet [1]", "bullet [2][3]"],
  "what_happened": ["para 1 [1]", "para 2"],
  "red_flags_intro": "Optional lead-in sentence for Red Flags to Watch.",
  "red_flags": ["flag [1]", "flag"],
  "recommended_actions": [
    {"audience": "All employees", "items": ["do this", "do that"]},
    {"audience": "IT Service Desk", "target": "within 30 days",
     "items": ["single action [2]"]}
  ],
  "questions": [["SOC", "question text"]],   # legacy alternative to recommended_actions
  "abbreviations": [{"term": "HHS OCR",
                     "definition": "Office for Civil Rights, U.S. Department of Health and Human Services",
                     "url": "https://..."}],   # optional glossary above Sources
  "sources": [
    {"text": "BleepingComputer: Headline (August 17, 2026)", "url": "https://..."},
    "Unlinked source, plain string still works"
  ],
  "footer_contact": "secops@genelabs.com  |  ServiceNow",
  "tlp": "AMBER+STRICT"                    # optional, defaults to AMBER+STRICT
}

Write inline citations in the content as bracketed numbers, e.g. "... in the
wild [1][3]." When the source is a paginated document, cite the page:
"[1, p. 5]", "[1, pp. 22-23]". The builder strips the space in front of the group
and renders it as superscript hard against the preceding character. Numbering in
the Sources list itself stays at normal size.
"""
import datetime, json, os, re, sys
from docx import Document
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.shared import Pt, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

# ---- GeneLabs corporate palette -------------------------------------------
ORANGE      = "FF7109"   # GeneLabs orange, sampled from the corporate deck rule
ORANGE_DEEP = "E05F00"   # deepened orange for small text
ABBEY       = "444648"   # GeneLabs Abbey, body and headline text
GREY        = "676765"   # corporate grey, wordmark and footer meta
RULE        = "D9D9D9"   # light grey: Word keeps it light on a dark canvas (verified)
LABEL_FILL  = "F4F4F4"
CALLOUT_FILL= "FFF4E8"
LINK_BLUE   = "1565C0"
FONT        = "Arial"

# Published bulletin category legend (unchanged)
CATEGORY_COLORS = {
    "THREAT BULLETIN": "C62828", "VULNERABILITY ALERT": "B71C1C",
    "MALWARE CAMPAIGN": "EF6C00", "THREAT ACTOR ACTIVITY": "E65100",
    "SUPPLY CHAIN ALERT": "FF8F00", "PEER INCIDENT": "6A1B9A",
    "AWARENESS BULLETIN": "1565C0", "INDUSTRY ALERT": "00796B",
}
SEVERITY_COLORS = {
    "CRITICAL": "B71C1C", "HIGH": "C62828", "MEDIUM": "EF6C00",
    "LOW": "2E7D32", "INFORMATIONAL": "1565C0",
}
DEFAULT_TLP = "AMBER+STRICT"
SERVICENOW_URL = ("https://<your-instance>.service-now.com/esp?id=sc_cat_item"
                  "&sys_id=2498318b4fee53c0f628d0af0310c75d")
# Matches [1], [1][2] and page-level cites such as [1, p. 5] or [1, pp. 22-23]
_ONE_CITE = r"\[\d+(?:,\s*pp?\.\s*\d+(?:\s*[\u2013-]\s*\d+)?(?:,\s*\d+)*)?\]"
CITE_RE = re.compile(r"\s*((?:" + _ONE_CITE + r")+)")


def _copyright(date_str):
    m = re.search(r"(20\d{2})", date_str or "")
    year = m.group(1) if m else str(datetime.date.today().year)
    return f"© {year} GeneLabs LLC All rights reserved."


# ---- low level helpers -----------------------------------------------------
def _run(p, text, size=10, bold=False, italic=False, color=ABBEY, caps=False,
         superscript=False):
    r = p.add_run(text)
    r.font.name = FONT
    r.font.size = Pt(size)
    r.font.bold = bold
    r.font.italic = italic
    r.font.all_caps = caps
    if superscript:
        r.font.superscript = True
    if color:
        r.font.color.rgb = RGBColor.from_string(color)
    rpr = r._element.get_or_add_rPr()
    rf = rpr.find(qn('w:rFonts'))
    if rf is None:
        rf = OxmlElement('w:rFonts'); rpr.insert(0, rf)
    for a in ('w:ascii', 'w:hAnsi', 'w:cs'):
        rf.set(qn(a), FONT)
    return r


def _cite(p, text, size=10, color=ABBEY, bold=False):
    """Write text, rendering any [n] / [n][m] groups as superscript and linking
    the first use of each glossary term to the Abbreviations section."""
    pos = 0
    for m in CITE_RE.finditer(text):
        if m.start() > pos:
            _text(p, text[pos:m.start()], size=size, color=color, bold=bold)
        _run(p, m.group(1), size=size, color=color, superscript=True)
        pos = m.end()
    if pos < len(text):
        _text(p, text[pos:], size=size, color=color, bold=bold)
    return p


def _link(p, text, url, size=9, color=LINK_BLUE, bold=False):
    part = p.part
    r_id = part.relate_to(url, RT.HYPERLINK, is_external=True)
    h = OxmlElement('w:hyperlink')
    h.set(qn('r:id'), r_id)
    run = OxmlElement('w:r')
    rpr = OxmlElement('w:rPr')
    style = OxmlElement('w:rStyle'); style.set(qn('w:val'), 'Hyperlink')
    fonts = OxmlElement('w:rFonts')
    for a in ('w:ascii', 'w:hAnsi', 'w:cs'):
        fonts.set(qn(a), FONT)
    col = OxmlElement('w:color'); col.set(qn('w:val'), color)
    sz = OxmlElement('w:sz'); sz.set(qn('w:val'), str(int(size * 2)))
    szcs = OxmlElement('w:szCs'); szcs.set(qn('w:val'), str(int(size * 2)))
    for el in (style, fonts, col, sz, szcs):
        rpr.append(el)
    if bold:
        rpr.append(OxmlElement('w:b'))
    run.append(rpr)
    t = OxmlElement('w:t'); t.set(qn('xml:space'), 'preserve'); t.text = text
    run.append(t)
    h.append(run)
    p._p.append(h)
    return h


_BOOKMARK_ID = [1000]
ABBREV_BOOKMARK = "CTI_Abbreviations"
# First occurrence of each glossary term is linked to the Abbreviations section.
_GLOSSARY = {"terms": [], "linked": set()}


def _bookmark(p, name):
    bid = str(_BOOKMARK_ID[0]); _BOOKMARK_ID[0] += 1
    start = OxmlElement('w:bookmarkStart')
    start.set(qn('w:id'), bid); start.set(qn('w:name'), name)
    end = OxmlElement('w:bookmarkEnd'); end.set(qn('w:id'), bid)
    p._p.insert(0, start)
    p._p.append(end)
    return p


def _internal_link(p, anchor, text, size=10, color=LINK_BLUE, bold=False):
    h = OxmlElement('w:hyperlink')
    h.set(qn('w:anchor'), anchor)
    run = OxmlElement('w:r')
    rpr = OxmlElement('w:rPr')
    style = OxmlElement('w:rStyle'); style.set(qn('w:val'), 'Hyperlink')
    fonts = OxmlElement('w:rFonts')
    for a in ('w:ascii', 'w:hAnsi', 'w:cs'):
        fonts.set(qn(a), FONT)
    col = OxmlElement('w:color'); col.set(qn('w:val'), color)
    sz = OxmlElement('w:sz'); sz.set(qn('w:val'), str(int(size * 2)))
    szcs = OxmlElement('w:szCs'); szcs.set(qn('w:val'), str(int(size * 2)))
    for el in (style, fonts, col, sz, szcs):
        rpr.append(el)
    if bold:
        rpr.append(OxmlElement('w:b'))
    run.append(rpr)
    t = OxmlElement('w:t'); t.set(qn('xml:space'), 'preserve'); t.text = text
    run.append(t)
    h.append(run)
    p._p.append(h)
    return h


def _pageref(p, anchor, size=9, color=GREY):
    r = p.add_run(); r.font.name = FONT; r.font.size = Pt(size)
    r.font.color.rgb = RGBColor.from_string(color)
    begin = OxmlElement('w:fldChar'); begin.set(qn('w:fldCharType'), 'begin')
    instr = OxmlElement('w:instrText'); instr.set(qn('xml:space'), 'preserve')
    instr.text = " PAGEREF %s \\h " % anchor
    sep = OxmlElement('w:fldChar'); sep.set(qn('w:fldCharType'), 'separate')
    t = OxmlElement('w:t'); t.text = "1"
    end = OxmlElement('w:fldChar'); end.set(qn('w:fldCharType'), 'end')
    for el in (begin, instr, sep, t, end):
        r._element.append(el)
    return r


def _text(p, text, size=10, color=ABBEY, bold=False):
    """Plain text, except the first occurrence of each glossary term, which
    becomes an internal link to the Abbreviations section."""
    pending = [t for t in _GLOSSARY["terms"] if t not in _GLOSSARY["linked"]]
    hit = None
    for term in pending:
        m = re.search(r"\b" + re.escape(term) + r"\b", text)
        if m and (hit is None or m.start() < hit[1].start()):
            hit = (term, m)
    if hit is None:
        if text:
            _run(p, text, size=size, color=color, bold=bold)
        return p
    term, m = hit
    _GLOSSARY["linked"].add(term)
    if m.start():
        _run(p, text[:m.start()], size=size, color=color, bold=bold)
    _internal_link(p, ABBREV_BOOKMARK, text[m.start():m.end()], size=size, bold=bold)
    return _text(p, text[m.end():], size=size, color=color, bold=bold)


def _para(doc, text="", **kw):
    space_after = kw.pop("space_after", 6)
    space_before = kw.pop("space_before", 0)
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.space_before = Pt(space_before)
    p.paragraph_format.line_spacing = 1.08
    if text:
        _run(p, text, **kw)
    return p


def _shade_p(p, fill):
    pPr = p._p.get_or_add_pPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear'); shd.set(qn('w:color'), 'auto'); shd.set(qn('w:fill'), fill)
    pPr.append(shd)


def _border(p, edge, color, sz=6, space=0):
    pPr = p._p.get_or_add_pPr()
    bdr = pPr.find(qn('w:pBdr'))
    if bdr is None:
        bdr = OxmlElement('w:pBdr'); pPr.append(bdr)
    e = OxmlElement('w:' + edge)
    e.set(qn('w:val'), 'single'); e.set(qn('w:sz'), str(sz))
    e.set(qn('w:space'), str(space)); e.set(qn('w:color'), color)
    bdr.append(e)


def _cell_shade(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear'); shd.set(qn('w:color'), 'auto'); shd.set(qn('w:fill'), fill)
    tcPr.append(shd)


def _table_borders(table, color=RULE, sz=6, insideH=True, insideV=True):
    tblPr = table._tbl.tblPr
    borders = OxmlElement('w:tblBorders')
    edges = ['top', 'left', 'bottom', 'right']
    if insideH: edges.append('insideH')
    if insideV: edges.append('insideV')
    for edge in edges:
        e = OxmlElement('w:' + edge)
        e.set(qn('w:val'), 'single'); e.set(qn('w:sz'), str(sz))
        e.set(qn('w:space'), '0'); e.set(qn('w:color'), color)
        borders.append(e)
    tblPr.append(borders)


def _fixed_layout(table, total_width):
    tblPr = table._tbl.tblPr
    lay = OxmlElement('w:tblLayout'); lay.set(qn('w:type'), 'fixed')
    tblPr.append(lay)
    w = OxmlElement('w:tblW')
    w.set(qn('w:w'), str(int(total_width.twips))); w.set(qn('w:type'), 'dxa')
    tblPr.append(w)


def _set_grid(table, widths):
    grid = table._tbl.find(qn('w:tblGrid'))
    if grid is None:
        grid = OxmlElement('w:tblGrid'); table._tbl.insert(1, grid)
    for gc in list(grid):
        grid.remove(gc)
    for w in widths:
        gc = OxmlElement('w:gridCol'); gc.set(qn('w:w'), str(int(w.twips))); grid.append(gc)
    for row in table.rows:
        for cell, w in zip(row.cells, widths):
            cell.width = w


def _no_split(table):
    for row in table.rows:
        trPr = row._tr.get_or_add_trPr()
        trPr.append(OxmlElement('w:cantSplit'))


def _keep_next(p):
    pPr = p._p.get_or_add_pPr()
    pPr.append(OxmlElement('w:keepNext'))


def _field(p, instr):
    r = p.add_run(); r.font.name = FONT; r.font.size = Pt(8)
    r.font.color.rgb = RGBColor.from_string(GREY)
    fld = OxmlElement('w:fldChar'); fld.set(qn('w:fldCharType'), 'begin')
    it = OxmlElement('w:instrText'); it.set(qn('xml:space'), 'preserve'); it.text = instr
    sep = OxmlElement('w:fldChar'); sep.set(qn('w:fldCharType'), 'separate')
    t = OxmlElement('w:t'); t.text = "1"
    end = OxmlElement('w:fldChar'); end.set(qn('w:fldCharType'), 'end')
    for el in (fld, it, sep, t, end):
        r._element.append(el)
    return r


def _bullet(doc):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.left_indent = Inches(0.28)
    p.paragraph_format.first_line_indent = Inches(-0.16)
    p.paragraph_format.line_spacing = 1.08
    _run(p, "▪  ", size=9, bold=True, color=ORANGE)
    return p


# ---- body blocks -----------------------------------------------------------
def add_meta_line(doc, report_id, date):
    p = _para(doc, space_after=2)
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    _run(p, report_id, size=9, bold=True, color=ABBEY)
    _run(p, "  |  ", size=9, color=GREY)
    _run(p, date, size=9, color=GREY)
    return p


def add_category_tag(doc, category):
    """Category chip: a single shaded cell sized to the label, as published."""
    label = category.upper()
    fill = CATEGORY_COLORS.get(label, "555555")
    width = Inches(0.084 * len(label) + 0.30)
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.autofit = False
    _fixed_layout(table, width)
    _set_grid(table, (width,))
    _table_borders(table, color=fill, sz=2, insideH=False, insideV=False)
    cell = table.rows[0].cells[0]
    _cell_shade(cell, fill)
    tcPr = cell._tc.get_or_add_tcPr()
    tcPr.append(OxmlElement('w:noWrap'))
    mar = OxmlElement('w:tcMar')
    for side in ('left', 'right'):
        m = OxmlElement('w:' + side)
        m.set(qn('w:w'), '60'); m.set(qn('w:type'), 'dxa')
        mar.append(m)
    tcPr.append(mar)
    p = cell.paragraphs[0]
    p.text = ""
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    r = _run(p, " " + label + " ", size=8, bold=True, color="FFFFFF")
    rpr = r._element.get_or_add_rPr()
    sp = OxmlElement('w:spacing'); sp.set(qn('w:val'), '14')
    rpr.append(sp)
    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_before = Pt(0)
    spacer.paragraph_format.space_after = Pt(0)
    spacer.paragraph_format.line_spacing = 1.0
    _run(spacer, "", size=3)
    return table


def add_title(doc, title):
    p = _para(doc, title, size=16, bold=True, color=ABBEY, space_after=3)
    p.paragraph_format.line_spacing = 1.0
    return p


def add_audience(doc, audience):
    p = _para(doc, space_after=8)
    _run(p, "Audience:  ", size=9, bold=True, color=GREY, caps=True)
    _run(p, audience, size=9, color=GREY)
    _border(p, 'bottom', RULE, sz=6, space=4)
    return p


def add_relevance_box(doc, severity, relevance, supersedes=None):
    sev = severity.upper()
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.space_after = Pt(10)
    p.paragraph_format.left_indent = Inches(0.08)
    p.paragraph_format.line_spacing = 1.08
    _shade_p(p, CALLOUT_FILL)
    _border(p, 'left', ORANGE, sz=24, space=6)
    _run(p, "Relevance to GeneLabs:  ", size=10, bold=True, color=ABBEY)
    _run(p, sev, size=10, bold=True, color=SEVERITY_COLORS.get(sev, ABBEY))
    tail = "  —  " + relevance
    if supersedes:
        tail += f" This bulletin supersedes {supersedes}."
    _cite(p, tail, size=10, color=ABBEY)
    return p


def add_glossary_note(doc):
    """Small pointer under the callout so the reader knows a glossary exists."""
    p = _para(doc, space_before=0, space_after=8)
    _run(p, "Acronyms used in this bulletin are defined under ", size=8,
         italic=True, color=GREY)
    _internal_link(p, ABBREV_BOOKMARK, "Abbreviations", size=8)
    _run(p, ", page ", size=8, italic=True, color=GREY)
    _pageref(p, ABBREV_BOOKMARK, size=8)
    _run(p, ". First use of each term links there.", size=8, italic=True, color=GREY)
    return p


def add_section_heading(doc, text):
    p = _para(doc, space_before=12, space_after=4)
    r = _run(p, text, size=10, bold=True, color=ORANGE_DEEP, caps=True)
    rpr = r._element.get_or_add_rPr()
    sp = OxmlElement('w:spacing'); sp.set(qn('w:val'), '20')
    rpr.append(sp)
    _border(p, 'bottom', ORANGE, sz=12, space=3)
    _keep_next(p)
    return p


def add_bullets(doc, items):
    for it in items:
        p = _bullet(doc)
        _cite(p, it, size=10, color=ABBEY)


def add_detail_table(doc, rows):
    if not rows:
        return
    table = doc.add_table(rows=len(rows), cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.autofit = False
    _fixed_layout(table, Inches(6.7))
    _table_borders(table)
    for i, row in enumerate(rows):
        k, v = (list(row) + ["", ""])[:2]
        c0, c1 = table.rows[i].cells
        _cell_shade(c0, LABEL_FILL)
        for cell, text, bold in ((c0, k, True), (c1, v, False)):
            para = cell.paragraphs[0]
            para.text = ""
            para.paragraph_format.space_before = Pt(2)
            para.paragraph_format.space_after = Pt(2)
            _cite(para, text, size=9, color=ABBEY, bold=bold)
    _set_grid(table, (Inches(1.85), Inches(4.85)))
    _no_split(table)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def add_recommended_actions(doc, groups):
    for g in groups:
        if isinstance(g, (list, tuple)):
            g = {"audience": g[0], "items": list(g[1]) if isinstance(g[1], (list, tuple)) else [g[1]]}
        aud = g.get("audience", "")
        target = g.get("target")
        items = g.get("items", [])
        label = aud + (f" (target: {target})" if target else "")
        if len(items) == 1:
            p = _bullet(doc)
            p.paragraph_format.space_before = Pt(5)
            _run(p, label + ": ", size=10, bold=True, color=ABBEY)
            _cite(p, items[0], size=10, color=ABBEY)
        else:
            p = _para(doc, space_before=4, space_after=3)
            _run(p, label + ":", size=10, bold=True, color=ABBEY)
            _keep_next(p)
            for it in items:
                b = _bullet(doc)
                _cite(b, it, size=10, color=ABBEY)


def add_questions(doc, questions):
    """Group questions under one label per team.

    A team is named once, in first-appearance order, with all of its questions
    bulleted beneath it. A team holding a single question renders inline on one
    bullet, matching the published house bulletins. Accepts either the flat
    [[team, question], ...] pair form or the grouped
    [{"audience": ..., "items": [...]}, ...] form.
    """
    order, grouped = [], {}
    for entry in questions:
        if isinstance(entry, dict):
            aud = entry.get("audience", "")
            items = entry.get("items", [])
        else:
            aud, items = entry[0], [entry[1]]
        if aud not in grouped:
            order.append(aud)
            grouped[aud] = []
        grouped[aud].extend(items)
    # Style is decided ONCE for the whole section, not per team. If any team holds
    # more than one question, every team gets its own label line - otherwise a
    # single-question team renders as one more bullet under the previous team's
    # list and reads as if it belonged to that team.
    stacked = any(len(grouped[a]) > 1 for a in order)
    for aud in order:
        items = grouped[aud]
        if not stacked:
            p = _bullet(doc)
            p.paragraph_format.space_before = Pt(5)
            _run(p, aud + ": ", size=10, bold=True, color=ABBEY)
            _cite(p, items[0], size=10, color=ABBEY)
        else:
            p = _para(doc, space_before=4, space_after=3)
            _run(p, aud + ":", size=10, bold=True, color=ABBEY)
            _keep_next(p)
            for it in items:
                b = _bullet(doc)
                _cite(b, it, size=10, color=ABBEY)


def add_abbreviations(doc, entries):
    """Compact glossary for acronyms too long to expand inline."""
    for e in entries:
        if isinstance(e, (list, tuple)):
            e = {"term": e[0], "definition": e[1] if len(e) > 1 else "",
                 "url": e[2] if len(e) > 2 else None}
        term = e.get("term", "")
        definition = e.get("definition", "")
        url = e.get("url")
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(2)
        p.paragraph_format.left_indent = Inches(0.28)
        p.paragraph_format.first_line_indent = Inches(-0.28)
        if url:
            _link(p, term, url, size=9, bold=True)
        else:
            _run(p, term, size=9, bold=True, color=ABBEY)
        _run(p, "  \u2014  ", size=9, color=GREY)
        _cite(p, definition, size=9, color=GREY)


def add_sources(doc, sources):
    for i, s in enumerate(sources, 1):
        if isinstance(s, dict):
            text, url = s.get("text", ""), s.get("url")
        elif isinstance(s, (list, tuple)):
            text, url = (list(s) + [None])[:2]
        else:
            text, url = s, None
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(2)
        p.paragraph_format.left_indent = Inches(0.28)
        p.paragraph_format.first_line_indent = Inches(-0.28)
        _run(p, f"[{i}]  ", size=9, bold=True, color=ORANGE_DEEP)
        if url:
            _link(p, text, url, size=9)
        else:
            _run(p, text, size=9, color=GREY)


def add_contact(doc, contact):
    p = _para(doc, space_before=12, space_after=0)
    _shade_p(p, LABEL_FILL)
    _run(p, "  Questions:  ", size=9, bold=True, color=ABBEY)
    for i, part in enumerate(re.split(r"\s*\|\s*", contact)):
        if i:
            _run(p, "  |  ", size=9, color=GREY)
        if "@" in part:
            _link(p, part, "mailto:" + part, size=9)
        elif part.strip().lower() == "servicenow":
            _link(p, part, SERVICENOW_URL, size=9)
        else:
            _run(p, part, size=9, color=ABBEY)
    _run(p, "  ", size=9, color=ABBEY)
    return p


# ---- footer furniture ------------------------------------------------------
def build_footer(doc, report_id, wordmark_path, tlp, copyright_line):
    footer = doc.sections[0].footer
    for p in list(footer.paragraphs):
        p._element.getparent().remove(p._element)
    for t in list(footer.tables):
        t._tbl.getparent().remove(t._tbl)

    rule = footer.add_paragraph()
    rule.paragraph_format.space_before = Pt(0)
    rule.paragraph_format.space_after = Pt(4)
    _border(rule, 'top', ORANGE, sz=12, space=1)

    table = footer.add_table(rows=1, cols=3, width=Inches(6.9))
    table.autofit = False
    _fixed_layout(table, Inches(6.9))
    widths = (Inches(1.25), Inches(3.55), Inches(2.10))
    _set_grid(table, widths)
    cells = table.rows[0].cells

    p0 = cells[0].paragraphs[0]
    p0.paragraph_format.space_after = Pt(0)
    if os.path.exists(wordmark_path):
        p0.add_run().add_picture(wordmark_path, width=Inches(0.95))

    p1 = cells[1].paragraphs[0]
    p1.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p1.paragraph_format.space_after = Pt(0)
    _run(p1, "TLP: ", size=8, bold=True, italic=True, color=ABBEY)
    _run(p1, tlp, size=8, bold=True, italic=True, color=SEVERITY_COLORS["MEDIUM"])
    _run(p1, "  |  CTI Intelligence Bulletin", size=8, color=GREY)
    p1b = cells[1].add_paragraph()
    p1b.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p1b.paragraph_format.space_after = Pt(0)
    _run(p1b, copyright_line, size=7, color=GREY)

    p2 = cells[2].paragraphs[0]
    p2.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p2.paragraph_format.space_after = Pt(0)
    _run(p2, "DOC # ", size=8, bold=True, color=ABBEY)
    _run(p2, report_id, size=8, color=GREY)
    p2b = cells[2].add_paragraph()
    p2b.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p2b.paragraph_format.space_after = Pt(0)
    _run(p2b, "Page ", size=8, color=GREY)
    _field(p2b, " PAGE ")
    _run(p2b, " of ", size=8, color=GREY)
    _field(p2b, " NUMPAGES ")


def set_base_style(doc):
    st = doc.styles['Normal']
    st.font.name = FONT
    st.font.size = Pt(10)
    st.font.color.rgb = RGBColor.from_string(ABBEY)
    rpr = st.element.get_or_add_rPr()
    rf = rpr.find(qn('w:rFonts'))
    if rf is None:
        rf = OxmlElement('w:rFonts'); rpr.insert(0, rf)
    for a in ('w:ascii', 'w:hAnsi', 'w:cs', 'w:eastAsia'):
        rf.set(qn(a), FONT)


DEFAULT_WORDMARK = os.path.join(os.path.dirname(os.path.abspath(__file__)),
                                "..", "assets", "genelabs_wordmark_grey.png")


def build(spec, template_path, out_path, wordmark_path=None):
    wordmark_path = wordmark_path or DEFAULT_WORDMARK
    abbrevs = spec.get("abbreviations") or []
    _GLOSSARY["terms"] = sorted(
        [(a[0] if isinstance(a, (list, tuple)) else a.get("term", "")) for a in abbrevs],
        key=len, reverse=True)
    _GLOSSARY["linked"] = set()
    doc = Document(template_path)
    set_base_style(doc)
    add_meta_line(doc, spec["report_id"], spec["date"])
    add_category_tag(doc, spec["category"])
    add_title(doc, spec["title"])
    if spec.get("audience"):
        add_audience(doc, spec["audience"])
    add_relevance_box(doc, spec["severity"], spec["relevance"], spec.get("supersedes"))
    if abbrevs:
        add_glossary_note(doc)
    if spec.get("detail_table"):
        add_detail_table(doc, spec["detail_table"])
    if spec.get("why_this_matters"):
        add_section_heading(doc, "Why This Matters")
        add_bullets(doc, spec["why_this_matters"])
    add_section_heading(doc, "What Happened")
    for para in spec.get("what_happened", []):
        p = _para(doc, space_after=7)
        _cite(p, para, size=10, color=ABBEY)
    if spec.get("red_flags") or spec.get("red_flags_intro"):
        add_section_heading(doc, "Red Flags to Watch")
        if spec.get("red_flags_intro"):
            p = _para(doc, space_after=4)
            _cite(p, spec["red_flags_intro"], size=10, color=ABBEY)
            _keep_next(p)
        add_bullets(doc, spec.get("red_flags", []))
    if spec.get("recommended_actions"):
        add_section_heading(doc, "Recommended Actions")
        add_recommended_actions(doc, spec["recommended_actions"])
    if spec.get("questions"):
        add_section_heading(doc, "Questions to Ask")
        add_questions(doc, spec["questions"])
    if abbrevs:
        _bookmark(add_section_heading(doc, "Abbreviations"), ABBREV_BOOKMARK)
        add_abbreviations(doc, abbrevs)
    if spec.get("sources"):
        add_section_heading(doc, "Sources")
        add_sources(doc, spec["sources"])
    add_contact(doc, spec.get("footer_contact", "secops@genelabs.com  |  ServiceNow"))
    build_footer(doc, spec["report_id"], wordmark_path,
                 spec.get("tlp", DEFAULT_TLP), _copyright(spec.get("date", "")))
    cp = doc.core_properties
    cp.title = f'{spec["report_id"]} {spec["title"]}'
    cp.category = spec["category"]
    cp.comments = "GeneLabs Enterprise Information Security | Cyber Threat Intelligence"
    doc.save(out_path)
    return out_path


if __name__ == "__main__":
    if len(sys.argv) < 3:
        print(__doc__)
        sys.exit(1)
    spec = json.load(open(sys.argv[1], encoding="utf-8"))
    here = os.path.dirname(os.path.abspath(__file__))
    template = sys.argv[3] if len(sys.argv) > 3 else \
        os.path.join(here, "..", "assets", "bulletin_template.docx")
    wordmark = sys.argv[4] if len(sys.argv) > 4 else None
    print("wrote", build(spec, template, sys.argv[2], wordmark))
