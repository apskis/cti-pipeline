#!/usr/bin/env python3
"""Regression test for the 2026-08-24 package removals.

`test_spec.json` deliberately carries every key April removed, each marked
SHOULD NOT RENDER. A correct builder ignores all of them. Run after any change to
build_hunt.py:

    python3 scripts/test_build_hunt.py
"""
from __future__ import annotations

import json
import subprocess
import sys
import tempfile
from pathlib import Path

HERE = Path(__file__).resolve().parent
# every string April removed from the package on 2026-08-24, plus the test markers
FORBIDDEN = ["Status", "Start Date", "End Date", "Reviewed By", "Pending Review", "Proposed",
             "Execution Log", "Outcome & Classification", "Findings Classification Key",
             "Coverage & Gaps", "Investigation Result", "Overall Finding", "Coverage",
             "SHOULD NOT RENDER", "should not render"]


def main() -> int:
    import docx
    with tempfile.TemporaryDirectory() as tmp:
        out = Path(tmp) / "TH26-99.docx"
        subprocess.run([sys.executable, str(HERE / "build_hunt.py"),
                        str(HERE / "test_spec.json"), str(out)], check=True)
        d = docx.Document(str(out))
        paras = [p.text for p in d.paragraphs if p.text.strip()]
        text = "\n".join(paras) + "\n" + "\n".join(
            c.text for t in d.tables for r in t.rows for c in r.cells)

    failures: list[str] = []
    hits = [f for f in FORBIDDEN if f in text]
    if hits:
        failures.append(f"removed strings present: {hits}")
    meta = d.tables[0]
    if len(meta.columns) != 2:
        failures.append(f"metadata table has {len(meta.columns)} columns, must be 2")
    for t in d.tables:
        if t.rows[0].cells[0].text.strip() == "ATT&CK" and len(t.columns) != 4:
            failures.append(f"ATT&CK table has {len(t.columns)} columns, must be 4")
    for need in ("TH26-99", "CONTROL", "BASELINE", "in(ComputerName"):
        if need not in text:
            failures.append(f"missing expected content: {need}")

    if failures:
        print("FAIL"); [print(" -", f) for f in failures]; return 1
    print(f"PASS: {len(paras)} paragraphs, no removed section, metadata 2 columns, ATT&CK 4 columns")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
