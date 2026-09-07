#!/usr/bin/env python3
"""
Append a branded PROPOSED FOLLOW ON HUNTS table to a completed GeneLabs
Threat Hunt Report (.docx).

Usage:
    python3 append_proposed_hunts_table.py <report.docx> <spec.json>

The spec JSON must carry a top level "proposed_hunts" array. Each entry:
{
  "id":          "TH26-03A"        (assigned pivot ID, or "-" if not spawned),
  "title":       "Lab instrument VNC lateral movement",
  "hypothesis":  "one line ABLE style summary",
  "source":      "TH26-03 Finding 3",
  "telemetry":   "NetworkReceiveAcceptIP4 + ProcessRollup2 - available",
  "priority":    "High" | "Medium" | "Low",
  "disposition": "Executed as pivot hunt" | "Queued for review" |
                 "Blocked - telemetry gap" | "Handed to <team>" |
                 "Detection engineering instead"
}

Run AFTER build_hunt_report.py. Styling matches the GeneLabs house template
(orange section heading with rule, grey grid table, Arial). Safe to run on a
report that already lacks the section; running twice appends twice, so run once.
"""
import json
import sys

from docx import Document
from docx.shared import Pt, RGBColor, Inches
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

ORANGE = "E05F00"
ORANGE_RULE = "FF7109"
ABBEY = "444648"
HDR_FILL = "E05F00"
BORDER = "D9D9D9"
STRIPE = "FFF8F2"

DISPO_COLOR = {
    "executed": "1B7837",   # green - executed as pivot hunt
    "queued": "B26A00",     # amber - queued for review
    "blocked": "6A51A3",    # purple - telemetry / licensing gap
    "handed": "1565C0",     # blue - handed to another team
    "detection": "1565C0",  # blue - detection engineering instead
}


def _cell_shade(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:fill"), fill)
    tcPr.append(shd)


def _table_borders(table, color=BORDER, size="6"):
    tblPr = table._tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), size)
        el.set(qn("w:color"), color)
        borders.append(el)
    tblPr.append(borders)


def _set_grid(table, widths):
    grid = table._tbl.find(qn("w:tblGrid"))
    if grid is None:
        grid = OxmlElement("w:tblGrid")
        table._tbl.insert(1, grid)
    for w in widths:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(int(w.inches * 1440)))
        grid.append(col)
    for row in table.rows:
        for i, cell in enumerate(row.cells):
            cell.width = widths[i]


def _run(p, text, size=9, bold=False, color=ABBEY, caps=False):
    r = p.add_run(text)
    r.font.name = "Arial"
    r.font.size = Pt(size)
    r.font.bold = bold
    r.font.color.rgb = RGBColor.from_string(color)
    if caps:
        r.font.all_caps = True
    return r


def _para_border(p, edge, color, sz=12, space=3):
    pPr = p._p.get_or_add_pPr()
    pbdr = OxmlElement("w:pBdr")
    el = OxmlElement(f"w:{edge}")
    el.set(qn("w:val"), "single")
    el.set(qn("w:sz"), str(sz))
    el.set(qn("w:space"), str(space))
    el.set(qn("w:color"), color)
    pbdr.append(el)
    pPr.append(pbdr)


def _dispo_color(text):
    t = (text or "").lower()
    for key, color in DISPO_COLOR.items():
        if key in t:
            return color
    return ABBEY


def main():
    if len(sys.argv) != 3:
        sys.exit(__doc__)
    doc_path, spec_path = sys.argv[1], sys.argv[2]
    with open(spec_path, encoding="utf-8") as fh:
        spec = json.load(fh)
    hunts = spec.get("proposed_hunts") or []
    if not hunts:
        print("no proposed_hunts in spec - nothing appended")
        return

    doc = Document(doc_path)

    # Section heading, matching h_section in the house builder.
    hp = doc.add_paragraph()
    hp.paragraph_format.space_before = Pt(14)
    hp.paragraph_format.space_after = Pt(6)
    r = _run(hp, "PROPOSED FOLLOW ON HUNTS", size=11, bold=True, color=ORANGE, caps=True)
    r.font.name = "Arial"
    _para_border(hp, "bottom", ORANGE_RULE)

    lead = doc.add_paragraph()
    lead.paragraph_format.space_after = Pt(6)
    _run(lead, "Every follow on candidate this hunt generated, with its disposition. "
               "Executed pivots carry a parent letter ID and are filed as their own "
               "reports in the year archive.", size=9)

    headers = ["ID", "Proposed hunt", "Source", "Telemetry basis", "Priority", "Disposition"]
    widths = [Inches(0.7), Inches(2.15), Inches(0.9), Inches(1.35), Inches(0.6), Inches(1.2)]

    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    _table_borders(table)
    for j, h in enumerate(headers):
        c = table.rows[0].cells[j]
        _cell_shade(c, HDR_FILL)
        p = c.paragraphs[0]
        _run(p, h, size=8.5, bold=True, color="FFFFFF")
    for i, h in enumerate(hunts):
        cells = table.add_row().cells
        if i % 2 == 1:
            for c in cells:
                _cell_shade(c, STRIPE)
        _run(cells[0].paragraphs[0], h.get("id", "-"), size=8.5, bold=True)
        tp = cells[1].paragraphs[0]
        _run(tp, h.get("title", ""), size=8.5, bold=True)
        if h.get("hypothesis"):
            _run(tp, "\n" + h["hypothesis"], size=8)
        _run(cells[2].paragraphs[0], h.get("source", ""), size=8.5)
        _run(cells[3].paragraphs[0], h.get("telemetry", ""), size=8.5)
        _run(cells[4].paragraphs[0], h.get("priority", ""), size=8.5)
        dp = cells[5].paragraphs[0]
        _run(dp, h.get("disposition", ""), size=8.5, bold=True,
             color=_dispo_color(h.get("disposition")))
    _set_grid(table, widths)

    doc.save(doc_path)
    print(f"appended proposed follow on hunts table ({len(hunts)} rows) to {doc_path}")


if __name__ == "__main__":
    main()
