#!/usr/bin/env python3
"""GeneLabs Employee Awareness — BLOG format builder.

A second format for the AWR series, added 2026-08-27 at April's request: shorter,
illustrated, and laid out to be read rather than filed.

WHY THIS IS A SEPARATE BUILDER AND NOT A FLAG ON THE OLD ONE

`build_awareness.py` wraps the CTI bulletin builder, and it should, because the
analytical set has to look like one family. But a bulletin layout carries a category
tag, an audience line, a severity callout, a detail table and a numbered source list
before it reaches a single sentence the reader can use. That furniture is the point in
a bulletin and it is the problem in an all-employee post: it signals "compliance
document", and people scroll past compliance documents.

So this builder does NOT wrap the bulletin builder. It draws its own layout with the
same palette, the same typeface and the same footer furniture, so it is unmistakably
GeneLabs, but it is organised like a post: one big claim, an image, a short section,
an image, and a card of things to do.

WHAT IT KEEPS FROM THE ANALYTICAL FORMAT, DELIBERATELY

  * TLP:CLEAR in the footer, enforced. The sanitisation rule does not relax because
    the format got friendlier — this is still the document with no access control.
  * The three universal instructions: report it, reporting is never a waste of time,
    nobody is in trouble for reporting. `verify_awareness.py` still has to pass.
  * The wordmark, the orange rule, the DOC # and Page X of Y.

WORD BUDGET. Target 350-450 words of body copy. The long-form post for the same topic
ran 997. If a draft is over 500, cut a section rather than tightening sentences: the
format's whole value is that a reader finishes it.

Usage:
  python build_awareness_blog.py spec.json out.docx <genelabs-cti-bulletin skill dir>
"""
import json, os, sys, glob, subprocess, tempfile

from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

ORANGE      = "FF7109"
ORANGE_DEEP = "E05F00"
ABBEY       = "444648"
GREY        = "676765"
TINT        = "FFF4E8"
RULE        = "D9D9D9"

HERE = os.path.dirname(os.path.abspath(__file__))


# ------------------------------------------------------------------ primitives
def _rgb(h):
    return RGBColor(int(h[0:2], 16), int(h[2:4], 16), int(h[4:6], 16))


def _run(p, text, size=11, bold=False, italic=False, color=ABBEY, font="Arial"):
    r = p.add_run(text)
    r.font.name = font
    r.font.size = Pt(size)
    r.font.bold = bold
    r.font.italic = italic
    r.font.color.rgb = _rgb(color)
    rpr = r._element.get_or_add_rPr()
    rf = OxmlElement("w:rFonts")
    for a in ("w:ascii", "w:hAnsi", "w:cs"):
        rf.set(qn(a), font)
    rpr.append(rf)
    return r


def _para(doc, before=0, after=6, align=None):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(before)
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing = 1.18
    if align is not None:
        p.alignment = align
    return p


def _shade(el, fill):
    sh = OxmlElement("w:shd")
    sh.set(qn("w:val"), "clear")
    sh.set(qn("w:fill"), fill)
    el.append(sh)


def _rule(doc, colour=ORANGE, weight=18, after=10):
    p = _para(doc, after=after)
    pPr = p._p.get_or_add_pPr()
    b = OxmlElement("w:pBdr")
    bot = OxmlElement("w:bottom")
    bot.set(qn("w:val"), "single")
    bot.set(qn("w:sz"), str(weight))
    bot.set(qn("w:space"), "1")
    bot.set(qn("w:color"), colour)
    b.append(bot)
    pPr.append(b)
    return p


def _field(p, instr):
    for tag, txt in (("begin", None), (None, instr), ("end", None)):
        if tag:
            f = OxmlElement("w:fldChar")
            f.set(qn("w:fldCharType"), tag)
            r = OxmlElement("w:r"); r.append(f); p._p.append(r)
        else:
            it = OxmlElement("w:instrText")
            it.set(qn("xml:space"), "preserve")
            it.text = txt
            r = OxmlElement("w:r"); r.append(it); p._p.append(r)


# ---------------------------------------------------------------------- layout
def _footer(doc, skill_dir, report_id, tlp):
    """Same furniture as every other document in the set, with two AWR deviations:
    the TLP value is house black rather than orange (April, 26 Aug — orange reads as
    an alert on a document meant to be routine), and the label names this series."""
    sec = doc.sections[0]
    f = sec.footer
    for p in list(f.paragraphs):
        p._element.getparent().remove(p._element)

    line = f.add_paragraph()
    line.paragraph_format.space_before = Pt(2)
    line.paragraph_format.space_after = Pt(0)
    pPr = line._p.get_or_add_pPr()
    b = OxmlElement("w:pBdr"); top = OxmlElement("w:top")
    top.set(qn("w:val"), "single"); top.set(qn("w:sz"), "12")
    top.set(qn("w:space"), "4"); top.set(qn("w:color"), ORANGE)
    b.append(top); pPr.append(b)

    wm = glob.glob(os.path.join(skill_dir, "assets", "*wordmark*.png"))
    if not wm:
        wm = glob.glob(os.path.join(HERE, "assets", "*wordmark*.png"))
    if wm:
        _run(line, "", 8)
        line.runs[-1].add_picture(wm[0], width=Inches(0.62))
        _run(line, "   ", 8)
    # Label and value are SEPARATE runs, matching the analytical builders. The gate
    # locates the value run by its text in order to assert its colour, and a combined
    # "TLP: CLEAR" run defeats that check — which it duly did on the first build here.
    # Conforming is right: a new format should not be the reason a safety check stops
    # working, and the check is the one that keeps the marking honest.
    _run(line, "TLP: ", 7.5, color=GREY)
    _run(line, tlp, 7.5, color=ABBEY)
    _run(line, "  |  Employee Security Awareness  |  ", 7.5, color=GREY)
    _run(line, "© 2026 GeneLabs LLC", 7.5, color=GREY)

    p2 = f.add_paragraph()
    p2.paragraph_format.space_before = Pt(0)
    _run(p2, f"DOC # {report_id}    Page ", 7.5, color=GREY)
    _field(p2, " PAGE ")
    _run(p2, " of ", 7.5, color=GREY)
    _field(p2, " NUMPAGES ")


def _image(doc, path, width_in=5.9):
    p = _para(doc, before=2, after=2, align=WD_ALIGN_PARAGRAPH.CENTER)
    p.add_run().add_picture(path, width=Inches(width_in))
    return p


def _caption(doc, text):
    p = _para(doc, before=0, after=8, align=WD_ALIGN_PARAGRAPH.CENTER)
    _run(p, text, 8.5, italic=True, color=GREY)
    return p


def _card(doc, heading, items):
    """The actions card. A single-cell shaded table, because a shaded block is the
    one piece of layout a reader's eye returns to after skimming."""
    t = doc.add_table(rows=1, cols=1)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = t.rows[0].cells[0]
    _shade(cell._tc.get_or_add_tcPr(), TINT)
    cell.width = Inches(6.4)

    p0 = cell.paragraphs[0]
    p0.paragraph_format.space_before = Pt(8)
    p0.paragraph_format.space_after = Pt(4)
    _run(p0, heading, 12, bold=True, color=ORANGE_DEEP)
    for it in items:
        p = cell.add_paragraph()
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(5)
        p.paragraph_format.left_indent = Inches(0.16)
        p.paragraph_format.line_spacing = 1.15
        _run(p, "▪  ", 11, bold=True, color=ORANGE)
        _run(p, it, 10.5, color=ABBEY)
    cell.paragraphs[-1].paragraph_format.space_after = Pt(9)
    return t


def build(spec, out_docx, skill_dir, art_dir):
    tlp = spec.get("tlp", "CLEAR")
    if tlp != "CLEAR":
        sys.exit("employee awareness posts are TLP:CLEAR by definition; got %r" % tlp)
    rid = spec["report_id"]
    if not rid.startswith("AWR-"):
        sys.exit("report_id must be AWR-YYYY-MM-DD; got %r" % rid)

    doc = Document()
    for s in doc.sections:
        s.top_margin = Inches(0.7); s.bottom_margin = Inches(0.7)
        s.left_margin = Inches(0.9); s.right_margin = Inches(0.9)

    # Kicker + headline. No category tag, no audience line, no severity chip: this
    # reader does not need the taxonomy and every line of it delays the point.
    k = _para(doc, after=2)
    _run(k, spec.get("kicker", "SECURITY, IN PLAIN ENGLISH").upper(), 8.5,
         bold=True, color=ORANGE_DEEP)
    h = _para(doc, after=4)
    _run(h, spec["title"], 22, bold=True, color=ABBEY)
    if spec.get("standfirst"):
        sf = _para(doc, after=8)
        _run(sf, spec["standfirst"], 12.5, color=GREY)
    _rule(doc, ORANGE, 18, after=8)

    for block in spec["blocks"]:
        kind = block.get("type", "text")
        if kind == "image":
            _image(doc, os.path.join(art_dir, block["scene"] + ".png"),
                   block.get("width", 5.9))
            if block.get("caption"):
                _caption(doc, block["caption"])
        elif kind == "heading":
            p = _para(doc, before=8, after=3)
            _run(p, block["text"], 13.5, bold=True, color=ABBEY)
        elif kind == "text":
            for para in block["paragraphs"]:
                p = _para(doc, after=7)
                _run(p, para, 11, color=ABBEY)
        elif kind == "pull":
            p = _para(doc, before=6, after=8)
            pPr = p._p.get_or_add_pPr()
            b = OxmlElement("w:pBdr"); left = OxmlElement("w:left")
            left.set(qn("w:val"), "single"); left.set(qn("w:sz"), "24")
            left.set(qn("w:space"), "10"); left.set(qn("w:color"), ORANGE)
            b.append(left); pPr.append(b)
            p.paragraph_format.left_indent = Inches(0.18)
            _run(p, block["text"], 13, bold=True, color=ORANGE_DEEP)
        elif kind == "card":
            _card(doc, block["heading"], block["items"])
        else:
            sys.exit("unknown block type %r" % kind)

    if spec.get("sources"):
        _rule(doc, RULE, 6, after=4)
        p = _para(doc, after=2)
        _run(p, "Where this comes from", 9, bold=True, color=GREY)
        for s in spec["sources"]:
            sp = _para(doc, after=1)
            _run(sp, s, 8.5, color=GREY)
    if spec.get("footer_contact"):
        p = _para(doc, before=6)
        _run(p, "Questions or something to report:  ", 9.5, bold=True, color=ABBEY)
        _run(p, spec["footer_contact"], 9.5, color=GREY)

    _footer(doc, skill_dir, rid, tlp)
    doc.save(out_docx)

    words = sum(len(x.split()) for b in spec["blocks"] if b.get("type") == "text"
                for x in b["paragraphs"])
    words += sum(len(i.split()) for b in spec["blocks"] if b.get("type") == "card"
                 for i in b["items"])
    return {"out": out_docx, "body_words": words,
            "images": sum(1 for b in spec["blocks"] if b.get("type") == "image")}


if __name__ == "__main__":
    if len(sys.argv) < 4:
        sys.exit(__doc__)
    spec = json.load(open(sys.argv[1], encoding="utf-8"))
    art = sys.argv[4] if len(sys.argv) > 4 else tempfile.mkdtemp()
    sys.path.insert(0, HERE)
    import awareness_art
    for name in awareness_art.SCENES:
        awareness_art.render(name, art)
    st = build(spec, sys.argv[2], sys.argv[3], art)
    print(json.dumps(st, indent=1))
    if st["body_words"] > 500:
        print("WARNING: body is %d words; target is 350-450. Cut a section, do not "
              "tighten sentences." % st["body_words"], file=sys.stderr)
