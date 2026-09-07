#!/usr/bin/env python3
"""GeneLabs Employee Awareness post builder.

Wraps the plugin-managed `genelabs-cti-bulletin` builder rather than
reimplementing it, so the branding cannot drift. NEVER edit that skill — it is
plugin-managed. This file imports from it and post-processes the result.

Two house deviations for the employee series, both deliberate:
  1. The footer TLP value renders in HOUSE BLACK, regular weight — not the orange
     italic severity colour. April, 2026-08-26: the orange read as an alert on a
     document whose whole purpose is to be calm and routine.
  2. The footer label reads "Employee Security Awareness", not "CTI Intelligence
     Bulletin", because this is a different document class with a different reader.

Usage:
  python build_awareness.py spec.json out.docx <genelabs-cti-bulletin skill dir>
"""
import sys, os, importlib.util

ABBEY = "444648"   # GeneLabs Abbey — the house black used for body text


def _load_bulletin_builder(skill_dir):
    path = os.path.join(skill_dir, "scripts", "build_bulletin.py")
    if not os.path.exists(path):
        sys.exit("cannot find build_bulletin.py under %s" % skill_dir)
    spec = importlib.util.spec_from_file_location("build_bulletin", path)
    mod = importlib.util.module_from_spec(spec)
    sys.modules["build_bulletin"] = mod
    spec.loader.exec_module(mod)
    return mod


def _restyle_footer(docx_path, tlp_value):
    """Set the TLP value to house black, regular weight, and relabel the footer.

    Targeted by RUN TEXT rather than by index: the bulletin builder may reorder
    footer runs and an index-based edit would silently recolour the wrong thing.
    """
    from docx import Document
    from docx.shared import RGBColor

    doc = Document(docx_path)
    black = RGBColor.from_string(ABBEY)
    tlp_hits = 0
    label_hits = 0

    for section in doc.sections:
        for table in section.footer.tables:
            for row in table.rows:
                for cell in row.cells:
                    for para in cell.paragraphs:
                        for run in para.runs:
                            txt = run.text
                            if txt.strip() == tlp_value.strip():
                                run.font.color.rgb = black
                                run.font.bold = False
                                run.font.italic = False
                                tlp_hits += 1
                            elif txt.strip() == "TLP:":
                                run.font.italic = False
                            elif "TLP:" in txt:
                                run.font.italic = False
                            if "CTI Intelligence Bulletin" in txt:
                                run.text = txt.replace(
                                    "CTI Intelligence Bulletin",
                                    "Employee Security Awareness")
                                label_hits += 1
    doc.save(docx_path)
    return tlp_hits, label_hits


def main():
    if len(sys.argv) < 4:
        sys.exit("usage: build_awareness.py spec.json out.docx <cti-bulletin skill dir>")
    spec_path, out_path, skill_dir = sys.argv[1], sys.argv[2], sys.argv[3]

    import json
    spec = json.load(open(spec_path, encoding="utf-8"))

    # Enforce the series invariants here so a hand-written spec cannot break them.
    if spec.get("tlp") != "CLEAR":
        sys.exit("REFUSING TO BUILD: an employee awareness post must be tlp CLEAR, got %r"
                 % spec.get("tlp"))
    if not str(spec.get("report_id", "")).startswith("AWR-"):
        sys.exit("REFUSING TO BUILD: report_id must be AWR-YYYY-MM-DD, got %r"
                 % spec.get("report_id"))
    if "questions" in spec:
        sys.exit("REFUSING TO BUILD: employee posts close with recommended_actions, "
                 "not questions. Employees need instructions, not audit questions.")
    if not spec.get("recommended_actions"):
        sys.exit("REFUSING TO BUILD: recommended_actions is required — the post must "
                 "tell the reader what to do.")
    ra = spec["recommended_actions"]
    if ra and isinstance(ra[0], str):
        sys.exit("REFUSING TO BUILD: recommended_actions must be "
                 '[{"audience": ..., "items": [...]}] — the flat list-of-strings form '
                 "raises AttributeError inside the bulletin builder.")

    mod = _load_bulletin_builder(skill_dir)
    template = os.path.join(skill_dir, "assets", "bulletin_template.docx")
    # build() is the public entry point; the module has no main().
    mod.build(spec, template, out_path)

    tlp_hits, label_hits = _restyle_footer(out_path, spec["tlp"])
    print("restyled footer: TLP runs recoloured=%d, label replaced=%d" % (tlp_hits, label_hits))
    if tlp_hits == 0:
        print("WARNING: no footer run matched the TLP value — the orange may remain. "
              "Check whether the bulletin builder changed its footer run structure.")
    print("wrote %s" % out_path)


if __name__ == "__main__":
    main()
